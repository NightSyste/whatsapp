import sys
import os
import json
import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

if hasattr(sys.stdin, 'reconfigure'):
    sys.stdin.reconfigure(encoding='utf-8')
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
EXCEL_PATH = os.path.join(BASE_DIR, "WhatsApp_Kontakte.xlsx")

def save_chats_to_excel(chats):
    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "WhatsApp Chats"
        ws.views.sheetView[0].showGridLines = True

        headers = ["Nr.", "Typ", "Name", "Telefonnummer / Chat-ID", "Letzte Nachricht", "Zeitpunkt", "Ungelesen"]
        ws.append(headers)

        header_font = Font(name="Segoe UI", size=10, bold=True, color="E0E0E0")
        header_fill = PatternFill(start_color="2A2A2A", end_color="2A2A2A", fill_type="solid")
        thin_border = Border(
            left=Side(style="thin", color="404040"),
            right=Side(style="thin", color="404040"),
            top=Side(style="thin", color="404040"),
            bottom=Side(style="thin", color="404040")
        )

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
        ws.row_dimensions[1].height = 26

        data_font = Font(name="Segoe UI", size=10, color="202020")

        for idx, c in enumerate(chats, 2):
            ws.row_dimensions[idx].height = 22
            cid = idx - 1
            chat_type = "Gruppe" if c.get("isGroup") else "Person"
            name = str(c.get("name", "Unbekannt"))
            phone = str(c.get("phone") or c.get("id") or "")
            last_msg = str(c.get("lastMessageText", ""))
            
            # Format time
            ts = c.get("timestamp")
            if ts:
                try:
                    time_str = datetime.datetime.fromtimestamp(ts).strftime("%d.%m.%Y %H:%M")
                except Exception:
                    time_str = str(ts)
            else:
                time_str = "-"

            unread = str(c.get("unreadCount", 0))

            ws.cell(row=idx, column=1, value=cid).alignment = Alignment(horizontal="center", vertical="center")
            ws.cell(row=idx, column=2, value=chat_type).alignment = Alignment(horizontal="center", vertical="center")
            ws.cell(row=idx, column=3, value=name).alignment = Alignment(horizontal="left", vertical="center")
            ws.cell(row=idx, column=4, value=phone).alignment = Alignment(horizontal="left", vertical="center")
            ws.cell(row=idx, column=5, value=last_msg).alignment = Alignment(horizontal="left", vertical="center")
            ws.cell(row=idx, column=6, value=time_str).alignment = Alignment(horizontal="center", vertical="center")
            ws.cell(row=idx, column=7, value=unread).alignment = Alignment(horizontal="center", vertical="center")

            for col in range(1, 8):
                ws.cell(row=idx, column=col).border = thin_border

        widths = {"A": 6, "B": 12, "C": 26, "D": 26, "E": 45, "F": 20, "G": 12}
        for col, w in widths.items():
            ws.column_dimensions[col].width = w

        wb.save(EXCEL_PATH)
        return {"status": "success", "message": f"{len(chats)} Chats automatisch in Excel gespeichert."}
    except Exception as e:
        return {"status": "error", "message": str(e)}

if __name__ == "__main__":
    if len(sys.argv) > 1:
        cmd = sys.argv[1]
        if cmd == "save_chats":
            raw = sys.stdin.read()
            data = json.loads(raw)
            print(json.dumps(save_chats_to_excel(data)))
